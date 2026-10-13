import csv
import io
import json

import pymupdf
from openpyxl import load_workbook

from src.exports.csv_exporter import entities_csv, findings_csv
from src.exports.excel_exporter import comparison_workbook, inventory_workbook
from src.exports.json_exporter import comparison_json, report_json
from src.exports.pdf_exporter import comparison_pdf, report_pdf
from src.services.report_service import build_report_data, component_report


def test_report_data_uses_stored_results(settings, processed):
    report = build_report_data(settings, processed["A"]["id"])
    assert report["application"]["name"] == "AutoDoc AI"
    assert report["document"]["revision"] == "A"
    assert len(report["components"]) == 6
    assert any(f["rule_id"] == "duplicate_component_id" for f in report["findings"])
    assert "does not establish safety" in " ".join(report["limitations"])


def test_all_report_formats(settings, processed):
    report = build_report_data(settings, processed["A"]["id"])
    pdf = report_pdf(report)
    assert pdf.startswith(b"%PDF")
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        text = "".join(p.get_text() for p in doc)
    assert "Architecture Analysis Report" in text and "WheelSpeedProcessor" in text and "Known limitations" in text

    wb = load_workbook(io.BytesIO(inventory_workbook(report)))
    assert {"Summary", "Components", "Interfaces", "Ports", "Signals", "Dependencies", "Findings"} <= set(wb.sheetnames)
    assert wb["Components"].max_row == 7  # header + 6 components

    data = json.loads(report_json(report))
    assert data["schema"] == "autodoc-ai/report/v1" and data["document"]["title"]

    rows = list(csv.DictReader(io.StringIO(entities_csv(report["components"]).decode("utf-8-sig"))))
    assert len(rows) == 6 and rows[0]["page_number"]
    frows = list(csv.DictReader(io.StringIO(findings_csv(report["findings"]).decode("utf-8-sig"))))
    assert len(frows) == len(report["findings"])


def test_empty_report_is_handled(settings, sample_pdfs):
    from src.services import document_service

    rev = document_service.upload_revision(settings, "a.pdf", sample_pdfs["A"].read_bytes(), new_document_title="Empty")
    report = build_report_data(settings, rev["id"])
    assert "no extracted architecture data" in report["executive_summary"]
    assert report_pdf(report).startswith(b"%PDF")
    assert load_workbook(io.BytesIO(inventory_workbook(report)))["Components"]["A2"].value.startswith("No data")


def test_csv_neutralises_formula_injection():
    rows = [{"name": "=HYPERLINK(\"http://evil\")", "entity_type": "component", "description": "+cmd", "attributes": {}}]
    text = entities_csv(rows).decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+cmd" in text


def test_comparison_exports(settings, processed):
    from src.services import document_service

    result = document_service.compare(settings, processed["A"]["id"], processed["B"]["id"]).to_dict()
    assert json.loads(comparison_json(result))["counts"]["entities_added"] >= 7
    wb = load_workbook(io.BytesIO(comparison_workbook(result)))
    assert "Entity changes" in wb.sheetnames and wb["Entity changes"].max_row > 5
    assert comparison_pdf(result, "AutoDoc AI", "Rev A", "Rev B").startswith(b"%PDF")


def test_component_report(settings, processed):
    rep = component_report(settings, processed["B"]["id"], "BrakeTorqueCoordinator")
    assert rep["component"]["external_id"] == "C-03"
    assert {p["name"] for p in rep["ports"]} == {"RP_VehicleState", "PP_BrakeTorqueCmd", "RP_PedalRequest"}
    assert any(s["name"] == "Sig_BrakeTorqueRequest" for s in rep["signals"])
    assert any(f["name"].startswith("FF-01") for f in rep["flows"])
