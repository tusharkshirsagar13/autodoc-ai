"""Excel workbooks: architecture inventory and revision comparison."""

from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "1B2A41"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _safe(value):
    if isinstance(value, (list, dict)):
        if isinstance(value, dict):
            value = {k: v for k, v in value.items() if k != "occurrences"}
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        value = "'" + value  # avoid formula injection from document text
    return value


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, str, int]], rows: list[dict]):
    ws = wb.create_sheet(title[:31])
    for idx, (header, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(idx)].width = width
    for r, row in enumerate(rows, start=2):
        for c, (_, key, _) in enumerate(columns, start=1):
            value = key(row) if callable(key) else row.get(key)
            ws.cell(row=r, column=c, value=_safe(value)).alignment = WRAP
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if not rows:
        ws.cell(row=2, column=1, value="No data extracted for this category.")
    return ws


def inventory_workbook(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    doc = report["document"]
    rows = [("Application", f"{report['application']['name']} — {report['application']['subtitle']}"),
            ("Generated", report["generated_at"]), ("Document", doc["title"]), ("Revision", doc["revision"]),
            ("Source status", doc["source_status"]), ("File", doc["filename"]), ("Pages", doc["pages"]),
            ("SHA-256", doc["sha256"]), ("Executive summary", report["executive_summary"])]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=_safe(v)).alignment = WRAP
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    base = [("Page", "page_number", 7), ("Section", "section", 30), ("Method", "extraction_method", 16),
            ("Confidence", "confidence", 11), ("Review", "review_status", 12), ("Evidence", "evidence", 60)]
    _sheet(wb, "Components", [("ID", "external_id", 10), ("Component", "name", 30), ("Description", "description", 60),
                              ("Type", lambda r: r["attributes"].get("type", ""), 18)] + base, report["components"])
    _sheet(wb, "Interfaces", [("Interface", "name", 26), ("Type", lambda r: r["attributes"].get("interface_type", ""), 18),
                              ("Data elements", lambda r: ", ".join(r["attributes"].get("data_elements", [])), 50),
                              ("Description", "description", 50)] + base, report["interfaces"])
    _sheet(wb, "Ports", [("Component", "component", 26), ("Port", "name", 26),
                         ("Direction", lambda r: r["attributes"].get("direction", ""), 12),
                         ("Interface", lambda r: r["attributes"].get("interface", ""), 24)] + base, report["ports"])
    _sheet(wb, "Signals", [("Signal", "name", 28)] + [
        (label, (lambda k: lambda r: r["attributes"].get(k, ""))(key), 16)
        for label, key in (("Interface", "interface"), ("Data type", "data_type"), ("Unit", "unit"), ("Range", "range"),
                           ("Producer", "producer"), ("Consumer", "consumer"))] + base, report["signals"])
    _sheet(wb, "Dependencies", [("Source", "source_name", 26), ("Relation", "relation_type", 14), ("Target", "target_name", 26),
                                ("Via", "via", 24), ("Description", "description", 40)] + base, report["relationships"])
    _sheet(wb, "Functional flows", [("Flow", "name", 30), ("Steps", "description", 70)] + base, report["functional_flows"])
    _sheet(wb, "Findings", [("Code", "code", 8), ("Severity", "severity", 10), ("Category", "category", 22),
                            ("Summary", "summary", 60), ("Explanation", "explanation", 60), ("Evidence", "evidence", 50),
                            ("Pages", lambda r: ", ".join(map(str, r.get("pages", []))), 10),
                            ("Detection", "detection_method", 26), ("Evidence kind", "evidence_kind", 16),
                            ("Review", "review_status", 16), ("Reviewer comment", "reviewer_comment", 30)], report["findings"])
    ws_lim = wb.create_sheet("Limitations")
    for i, text in enumerate(report["limitations"], start=1):
        ws_lim.cell(row=i, column=1, value=text).alignment = WRAP
    ws_lim.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def comparison_workbook(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Comparison", f"Revision {result['old_label']} → {result['new_label']}"])
    for key, value in result["counts"].items():
        ws.append([key.replace("_", " ").capitalize(), value])
    ws.append(["Structural similarity", result["structural_similarity"]])
    for w in result["warnings"]:
        ws.append(["Warning", w])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    entity_rows = []
    for etype, bucket in result["entity_changes"].items():
        for kind in ("added", "removed", "changed"):
            for item in bucket[kind]:
                entity_rows.append({"type": etype, "change": kind, "name": item["name"], "component": item.get("component", ""),
                                    "page_old": item.get("page_old"), "page_new": item.get("page_new"),
                                    "details": "; ".join(f"{d['field']}: {d['old']} → {d['new']}" for d in item.get("differences", []))
                                    or item.get("description", "")})
    _sheet(wb, "Entity changes", [("Type", "type", 14), ("Change", "change", 10), ("Name", "name", 30),
                                  ("Component", "component", 22), ("Page (old)", "page_old", 10),
                                  ("Page (new)", "page_new", 10), ("Details", "details", 80)], entity_rows)
    section_rows = ([{"change": "added", "title": s["title"], "pages": s.get("pages_new"), "detail": s["text"][:2000]} for s in result["added_sections"]]
                    + [{"change": "removed", "title": s["title"], "pages": s.get("pages_old"), "detail": s["text"][:2000]} for s in result["removed_sections"]]
                    + [{"change": "modified", "title": s["title"], "pages": s.get("pages_new"),
                        "detail": "\n".join([f"+ {x}" for x in s["added_lines"]] + [f"- {x}" for x in s["removed_lines"]]
                                            + [f"~ {c['old']}  →  {c['new']}" for c in s["changed_lines"]])[:4000]}
                       for s in result["modified_sections"]])
    _sheet(wb, "Section changes", [("Change", "change", 10), ("Section", "title", 40),
                                   ("Pages", lambda r: ", ".join(map(str, r.get("pages") or [])), 10),
                                   ("Detail", "detail", 120)], section_rows)
    rel_rows = ([{**r, "change": "added"} for r in result["relationship_changes"].get("added", [])]
                + [{**r, "change": "removed"} for r in result["relationship_changes"].get("removed", [])])
    _sheet(wb, "Relationship changes", [("Change", "change", 10), ("Source", "source", 26), ("Relation", "relation", 14),
                                        ("Target", "target", 26), ("Via", "via", 24),
                                        ("Page", lambda r: r.get("page_new") or r.get("page_old"), 8)], rel_rows)
    _sheet(wb, "Affected dependencies", [("Source", "source", 26), ("Relation", "relation", 14), ("Target", "target", 26),
                                         ("Via", "via", 24), ("Present in", "present_in", 12), ("Changed items", "reasons", 70)],
           result["affected_dependencies"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
